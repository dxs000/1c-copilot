"""Синтетические фикстуры, повторяющие структуру реальных документов проекта (без реальных данных)."""

from __future__ import annotations

from email.message import EmailMessage as MimeMessage
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

THREAD_BODY = """External mail : Think before you Click !

Добрый день, коллеги.

Направляем на согласование ДС №10 на обновление конфигурации 1С УТ11 до версии 11.5.27.75.

С уважением,
Петрова Анна
Руководитель проекта
+7 (495) 111-22-33 #100
a.petrova@integrator.ru <mailto:a.petrova@integrator.ru>

From: Иванов Сергей Петрович
Sent: Tuesday, September 22, 2026 12:24 PM
To: Smith John <john.smith@customer.com>; Петрова Анна <a.petrova@integrator.ru>
Subject: Re: Обновление конфигурации 1С УТ11 до версии 11.5.27.75

Добрый день! Так как для УТ 11.5.27.75 минимальная версия платформы 8.3.27.1859, то 8.3.27.2342 подходит.

С уважением,
Сергей Иванов

________________________________

From: Smith John [mailto:john.smith@customer.com]

Sent: Tuesday, September 22, 2026 at 10:37 AM

To: Петрова Анна <a.petrova@integrator.ru> <mailto:a.petrova@integrator.ru>

Subject: Обновление конфигурации 1С УТ11 до версии 11.5.27.75

\tУважаемые коллеги,

\tВ четверг планируем обновить платформу до версии 8.3.27.2342 на тестовом окружении (сервер custsrvapp01).
\tПрошу подтвердить совместимость.

\tС уважением Джон

\tConformément à la Charte de la Déconnexion, les emails envoyés le soir n'appellent pas de réponse immédiate.
\tNotice:
\tThis message and any attachments are confidential and intended solely for the named recipients.
"""


def make_eml(path: Path, attachments: list[tuple[str, bytes, str]] = ()) -> Path:
    m = MimeMessage()
    m["Subject"] = "RE: Обновление конфигурации 1С УТ11 до версии 11.5.27.75"
    m["From"] = "Петрова Анна <a.petrova@integrator.ru>"
    m["To"] = "Smith John <john.smith@customer.com>"
    m["Date"] = "Wed, 23 Sep 2026 19:02:39 +0300"
    m.set_content(THREAD_BODY)
    for name, data, mime in attachments:
        maintype, subtype = mime.split("/")
        m.add_attachment(data, maintype=maintype, subtype=subtype, filename=name)
    path.write_bytes(bytes(m))
    return path


def _outline(paragraph, level: int) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    el = OxmlElement("w:outlineLvl")
    el.set(qn("w:val"), str(level))
    ppr.append(el)


def _strike(run) -> None:
    rpr = run._r.get_or_add_rPr()
    rpr.append(OxmlElement("w:strike"))


def make_tz(path: Path) -> Path:
    """ТЗ: заголовки через outlineLvl в стиле Normal, план тестирования с объединёнными ячейками,
    таблица покрытия, зачёркнутый текст, комментарий рецензента."""
    doc = Document()
    for text in ("Техническое задание", "на обновление конфигурации «1С:Управление торговлей» с версии 11.5.19.55 "
                                         "до версии 11.5.27.75"):
        doc.add_paragraph().add_run(text).bold = True
    doc.add_paragraph("Версия документа 2.0")
    toc = doc.add_paragraph("1. Общие положения\t1")
    toc.style = doc.styles.add_style("toc 1", 1)
    p = doc.add_paragraph()
    p.add_run("1. Общие положения").bold = True
    _outline(p, 0)
    p = doc.add_paragraph()
    p.add_run("1.3. Объем работ Исполнителя").bold = True
    _outline(p, 1)
    para = doc.add_paragraph("Исполнитель переносит доработки, включая справочник «Предопределенные значения (КС)». ")
    _strike(para.add_run("Удалённая формулировка про закрытие месяца."))
    doc.add_comment(para.runs[0], text="Не требуется, не актуально.", author="Рецензент")

    p = doc.add_paragraph("Приложение 1")
    _outline(p, 0)
    p = doc.add_paragraph()
    p.add_run("План тестирования").bold = True
    _outline(p, 1)
    t = doc.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    for cell, text in zip(t.rows[0].cells, ["№п/п", "Наименование объекта", "Процедура проверки", ""], strict=True):
        cell.text = text
    t.rows[0].cells[2].merge(t.rows[0].cells[3])
    group = t.add_row().cells
    group[0].merge(group[3]).text = "3. НСИ"
    first = t.add_row().cells
    first[0].text, first[1].text = "64", "Номенклатура"
    first[2].merge(first[3]).text = "В карточке элемента сверяем все реквизиты"
    second = t.add_row().cells
    second[2].merge(second[3]).text = "Переходим в Штрихкоды номенклатуры, реквизит «Номенклатура (КС)»"
    first[0].merge(second[0])  # вертикальное объединение №
    first[1].merge(second[1])
    third = t.add_row().cells
    third[0].text, third[1].text = "65", "Номенклатура контрагентов"
    third[2].merge(third[3]).text = "Сверяем реквизитный состав"

    p = doc.add_paragraph("Приложение 2")
    _outline(p, 0)
    cov = doc.add_table(rows=3, cols=4)
    cov.style = "Table Grid"
    for row, values in zip(cov.rows, [["№", "Документ, версия, дата", "Пункты Плана тестирования", "Покрытие Планом"],
                                      ["1.", "ПиМИ «01 НСИ», версия 1.0", "64, 70–72", "Частичное"],
                                      ["2.", "ПиМИ «Склад», версия 1.0", "12", "Полное"]], strict=True):
        for cell, text in zip(row.cells, values, strict=True):
            cell.text = text
    doc.save(path)
    return path


def make_pimi(path: Path) -> Path:
    """ПиМИ: стиль заголовка на базе Heading 1 с многоуровневым списком, номера автонумерации в тексте
    отсутствуют; тест-кейс разнесён на несколько строк с одинаковым №."""
    doc = Document()
    doc.add_paragraph().add_run("ПРОГРАММА и МЕТОДИКА ИСПЫТАНИЙ").bold = True
    doc.add_paragraph().add_run("Выгрузка и загрузка НСИ").bold = True
    h1 = doc.styles.add_style("1. Стиль 1 ГФ", 1)
    h1.base_style = doc.styles["Heading 1"]

    def numbered(text: str, ilvl: int):
        p = doc.add_paragraph(text, style=h1)
        num_pr = OxmlElement("w:numPr")
        lvl = OxmlElement("w:ilvl")
        lvl.set(qn("w:val"), str(ilvl))
        num_id = OxmlElement("w:numId")
        num_id.set(qn("w:val"), "1")
        num_pr.append(lvl)
        num_pr.append(num_id)
        p._p.get_or_add_pPr().append(num_pr)
        return p

    numbered("Общие сведения", 0)
    doc.add_paragraph("Обработка ПФ.ВыгрузкаНСИ.epf отбирает номенклатуру за последние 36 месяцев.")
    numbered("Программы испытаний", 0)
    numbered("Программа испытаний справочника \"Номенклатура\"", 1)
    numbered("Регистрация номенклатуры к выгрузке", 2)
    t = doc.add_table(rows=1, cols=5)
    t.style = "Table Grid"
    for cell, text in zip(t.rows[0].cells, ["№ п/п", "Функция", "Методика проверки", "Критерий успешности",
                                            "Результат"], strict=True):
        cell.text = text
    g = t.add_row().cells
    g[0].merge(g[4]).text = "Работа с обработкой. Выборка отдельных элементов номенклатуры"
    for method, criterion, result in (("Указать группу выгружаемых данных", "Правила подобраны", "Работает"),
                                      ("Установить отбор по группе", "ТЧ \"Номенклатура\" заполнена", "Работает"),
                                      ("Проверить ШтрихкодыНоменклатуры", "Штрихкоды перенесены", "Не работает")):
        cells = t.add_row().cells
        for cell, text in zip(cells, ["1", "Выбор группы данных", method, criterion, result], strict=True):
            cell.text = text
    cells = t.add_row().cells
    for cell, text in zip(cells, ["2", "Артикул", "Открыть карточку", "Артикул и код совпадают", "Работает"],
                          strict=True):
        cell.text = text
    numbered("Программа испытаний справочника \"Контрагенты\"", 1)
    doc.save(path)
    return path
