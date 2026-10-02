"""Разбор .docx: разделы по стилям заголовков, таблицы, тест-кейсы ПиМИ.

Чанк ТЗ/ДС — раздел документа (по заголовку). Чанк ПиМИ — один тест-кейс вместе с
названием своего раздела: строки-заголовки таблицы с объединёнными ячейками задают раздел.
"""

from __future__ import annotations

import re
from collections.abc import Iterable, Iterator
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.document import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph

from copilot1c.models import Chunk, DocType

_HEADING_RE = re.compile(r"^(Heading|Заголовок)\s*(\d+)$", re.IGNORECASE)
# Составные идентификаторы метаданных: ШтрихкодыНоменклатуры, ВыгрузкаНСИ, КС_Гамма
_MD_NAME_RE = re.compile(r"\bКС_[\wЁё]+|\b[А-ЯЁ][а-яё]+(?:[А-ЯЁ]+[а-яё]*)+\b")

# Ключевые слова заголовков колонок таблицы тест-кейсов
_COLS = {
    "num": ("№", "n п/п", "номер"),
    "function": ("функция", "проверяемая", "наименование"),
    "method": ("методика", "порядок", "шаги"),
    "criterion": ("критерий", "ожидаемый"),
    "result": ("результат", "статус", "отметка"),
}


@dataclass
class Section:
    path: list[str]  # иерархия заголовков
    paragraphs: list[str] = field(default_factory=list)

    @property
    def title(self) -> str:
        return " / ".join(self.path) or "(без заголовка)"

    @property
    def number(self) -> str:
        m = re.match(r"^\s*(\d+(?:\.\d+)*)", self.path[-1]) if self.path else None
        return m.group(1) if m else ""


@dataclass
class TestCase:
    num: str
    section: str
    function: str = ""
    method: str = ""
    criterion: str = ""
    result: str = ""
    objects: list[str] = field(default_factory=list)

    def text(self) -> str:
        parts = [f"Тест-кейс № {self.num}. Раздел: {self.section}"]
        for label, value in (
            ("Функция", self.function),
            ("Методика проверки", self.method),
            ("Критерий успешности", self.criterion),
            ("Результат", self.result),
        ):
            if value:
                parts.append(f"{label}: {value}")
        return "\n".join(parts)


def find_objects(text: str, known: Iterable[str] | None = None) -> list[str]:
    """Объекты метаданных в тексте: по реестру known (если задан) и по форме идентификатора."""
    found = {m.group(0) for m in _MD_NAME_RE.finditer(text)}
    if known:
        for name in known:
            if re.search(rf"(?<![\wЁё]){re.escape(name)}(?![\wЁё])", text):
                found.add(name)
    return sorted(found)


def iter_blocks(doc: DocxDocument) -> Iterator[Paragraph | Table]:
    """Абзацы и таблицы в порядке следования в документе."""
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            yield Paragraph(child, doc)
        elif tag == "tbl":
            yield Table(child, doc)


def _heading_level(p: Paragraph) -> int | None:
    m = _HEADING_RE.match(p.style.name if p.style is not None else "")
    return int(m.group(2)) if m else None


def _row_cells(row) -> list[str]:
    # Объединённые ячейки python-docx возвращает повторно — схлопываем подряд идущие дубли
    out: list[str] = []
    prev = None
    for cell in row.cells:
        if cell._tc is prev:
            continue
        prev = cell._tc
        out.append(cell.text.strip())
    return out


def _table_text(table: Table) -> str:
    return "\n".join(" | ".join(_row_cells(r)) for r in table.rows)


def _match_columns(header: list[str]) -> dict[str, int] | None:
    mapping: dict[str, int] = {}
    for i, cell in enumerate(header):
        low = cell.casefold()
        for key, words in _COLS.items():
            if key not in mapping and any(w in low for w in words):
                mapping[key] = i
                break
    if "num" in mapping and ("method" in mapping or "criterion" in mapping):
        return mapping
    return None


def table_test_cases(table: Table, section: str, known_objects: Iterable[str] | None = None) -> list[TestCase]:
    rows = [_row_cells(r) for r in table.rows]
    if not rows:
        return []
    cols = _match_columns(rows[0])
    if cols is None:
        return []
    cases: list[TestCase] = []
    current = section
    for cells in rows[1:]:
        non_empty = {c for c in cells if c}
        if len(non_empty) == 1 and (len(cells) == 1 or len(set(cells)) == 1 or not cells[cols["num"]]):
            current = f"{section} / {non_empty.pop()}"  # строка-заголовок подраздела
            continue
        def get(k: str, cells: list[str] = cells) -> str:
            return cells[cols[k]] if k in cols and cols[k] < len(cells) else ""

        if not get("num"):
            continue
        tc = TestCase(
            num=get("num").rstrip("."),
            section=current,
            function=get("function"),
            method=get("method"),
            criterion=get("criterion"),
            result=get("result"),
        )
        tc.objects = find_objects(tc.text(), known_objects)
        cases.append(tc)
    return cases


@dataclass
class ParsedDocx:
    title: str
    sections: list[Section]
    test_cases: list[TestCase]


def parse_docx(path: str | Path, known_objects: Iterable[str] | None = None) -> ParsedDocx:
    doc = Document(str(path))
    title = doc.core_properties.title or Path(path).stem
    sections: list[Section] = [Section(path=[])]
    test_cases: list[TestCase] = []
    stack: list[str] = []
    for block in iter_blocks(doc):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            level = _heading_level(block)
            if level and text:
                stack = stack[: level - 1] + [text]
                sections.append(Section(path=list(stack)))
            elif text:
                sections[-1].paragraphs.append(text)
        else:
            cases = table_test_cases(block, sections[-1].title, known_objects)
            if cases:
                test_cases.extend(cases)
            else:
                sections[-1].paragraphs.append(_table_text(block))
    sections = [s for s in sections if s.paragraphs]
    return ParsedDocx(title=title, sections=sections, test_cases=test_cases)


def guess_doc_type(filename: str) -> DocType:
    low = filename.casefold()
    if "пими" in low or "пми" in low or "методик" in low:
        return DocType.PIMI
    if re.search(r"(^|[^а-яё])тз([^а-яё]|$)", low) or "техническое задание" in low:
        return DocType.TZ
    if re.search(r"(^|[^а-яё])дс([^а-яё]|$)", low) or "соглашени" in low:
        return DocType.DS
    return DocType.DOC


def guess_doc_version(filename: str) -> str | None:
    m = re.search(r"(?:ред\.?|ver\.?|v)\s*(\d+)", filename, re.IGNORECASE)
    return m.group(1) if m else None


def docx_chunks(
    path: str | Path, project: str = "", known_objects: Iterable[str] | None = None,
    source: str | None = None,
) -> list[Chunk]:
    path = Path(path)
    parsed = parse_docx(path, known_objects)
    doc_type = guess_doc_type(path.name)
    version = guess_doc_version(path.name)
    src = source or str(path)
    chunks: list[Chunk] = []
    for s in parsed.sections:
        body = "\n".join(s.paragraphs)
        chunks.append(
            Chunk(
                text=f"{parsed.title}\n{s.title}\n\n{body}",
                doc_type=doc_type, source=src, title=s.title, project=project,
                doc_version=version, objects=find_objects(body, known_objects),
                extra={"section": s.number} if s.number else {},
            )
        )
    for tc in parsed.test_cases:
        chunks.append(
            Chunk(
                text=f"{parsed.title}\n{tc.text()}",
                doc_type=DocType.PIMI, source=src, title=f"Тест-кейс № {tc.num}", project=project,
                doc_version=version, objects=tc.objects,
                extra={"test_case": tc.num, "result": tc.result},
            )
        )
    return chunks
